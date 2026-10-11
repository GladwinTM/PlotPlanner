"""Document upload and processing.

Flow: validate type by content (not just the extension) and size, store the
bytes privately in Postgres, extract text, then ask the AI for facts. Facts
are saved *unconfirmed* with a link to the source document; the user reviews
them before they count as part of the profile.

Uploaded text is untrusted: it is passed to the model inside <document> tags
under a prompt that forbids following its instructions, and document
processing has no tool access.
"""
import asyncio
import io
import logging
import os
import shutil
import zipfile
from datetime import UTC, datetime
from uuid import UUID

from sqlalchemy import delete, func, literal_column, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import undefer

from app.ai import prompts
from app.ai.llm_client import LLMClient
from app.ai.schemas import DocumentExtraction
from app.core.exceptions import AppError, BadRequestError, NotFoundError, PayloadTooLargeError
from app.db.models import Document, ProfileFact

logger = logging.getLogger(__name__)

MAX_DOCUMENTS = 50
MAX_EXTRACTED_CHARS = 200_000
AI_CHARS = 24_000
OCR_MAX_PAGES = 10
OCR_SCALE = 2  # 144 DPI: legible for resumes without unbounded raster memory.
ALLOWED = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
}


def detect_type(filename: str, data: bytes) -> str:
    """Return the canonical content type, or raise if the bytes don't match an allowed format."""
    extension = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if extension not in ALLOWED:
        raise BadRequestError("Only PDF, DOCX and TXT files are supported", code="unsupported_file_type")
    if extension == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise BadRequestError("This file is not a valid PDF", code="invalid_file_content")
    elif extension == ".docx":
        if not data.startswith(b"PK\x03\x04"):
            raise BadRequestError("This file is not a valid DOCX document", code="invalid_file_content")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise BadRequestError("This file is not a valid DOCX document", code="invalid_file_content")
        except zipfile.BadZipFile as exc:
            raise BadRequestError("This file is not a valid DOCX document", code="invalid_file_content") from exc
    else:
        if b"\x00" in data[:8192]:
            raise BadRequestError("This file is not a plain text file", code="invalid_file_content")
        try:
            data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise BadRequestError("Text files must be UTF-8 encoded", code="invalid_file_content") from exc
    return ALLOWED[extension]


def _extract_pdf(data: bytes) -> str:
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as document:
        if document.needs_pass:
            raise BadRequestError("Password-protected PDFs are not supported", code="invalid_file_content")
        text = "\n".join(page.get_text() for page in document).strip()
        if text:
            return text
        if len(document) > OCR_MAX_PAGES:
            raise BadRequestError(
                f"Scanned PDFs can contain at most {OCR_MAX_PAGES} pages.", code="ocr_page_limit",
            )
        return _ocr_pdf(document)


def _ocr_pdf(document) -> str:
    """OCR an image-only PDF with bounded, local raster work."""
    try:
        import pymupdf
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise BadRequestError(
            "This scanned PDF needs OCR, but OCR is not available on this server yet.", code="ocr_unavailable",
        ) from exc
    command = os.getenv("TESSERACT_CMD") or shutil.which("tesseract") or "/opt/homebrew/opt/tesseract/bin/tesseract"
    if not os.path.isfile(command):
        raise BadRequestError(
            "This scanned PDF needs OCR, but the Tesseract engine is not available on this server yet.", code="ocr_unavailable",
        )
    pytesseract.pytesseract.tesseract_cmd = command
    pages: list[str] = []
    for page in document:
        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(OCR_SCALE, OCR_SCALE), alpha=False)
        image = Image.open(io.BytesIO(pixmap.tobytes("png")))
        pages.append(pytesseract.image_to_string(image, lang="eng"))
    return "\n".join(pages).strip()


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(content_type: str, data: bytes) -> str:
    if content_type == "application/pdf":
        text = _extract_pdf(data)
    elif content_type == ALLOWED[".docx"]:
        text = _extract_docx(data)
    else:
        text = data.decode("utf-8")
    return text.replace("\x00", "").strip()[:MAX_EXTRACTED_CHARS]


async def upload(session: AsyncSession, llm: LLMClient, user_id: UUID, filename: str, data: bytes, max_bytes: int) -> Document:
    if not data:
        raise BadRequestError("The file is empty", code="empty_file")
    if len(data) > max_bytes:
        raise PayloadTooLargeError(f"Files must be at most {max_bytes // (1024 * 1024)} MB")
    safe_name = (filename or "document").replace("/", "_").replace("\\", "_")[:255]
    content_type = detect_type(safe_name, data)
    count = await session.scalar(select(func.count()).select_from(Document).where(Document.user_id == user_id))
    if count >= MAX_DOCUMENTS:
        raise BadRequestError(f"You can store at most {MAX_DOCUMENTS} documents. Delete one to upload another.")

    document = Document(
        user_id=user_id, filename=safe_name, content_type=content_type, size_bytes=len(data), content=data,
        processing_status="processing",
    )
    session.add(document)
    await session.commit()
    return await _process(session, llm, document)


async def _process(session: AsyncSession, llm: LLMClient, document: Document) -> Document:
    try:
        text = await asyncio.to_thread(extract_text, document.content_type, document.content)
    except AppError as exc:
        return await _fail(session, document, exc.message)
    except Exception:
        logger.warning("text extraction failed for a %s upload", content_type)
        return await _fail(session, document, "Text could not be extracted from this file.")
    if not text:
        return await _fail(session, document, "No readable text was found. Try a clearer scan or a text-based PDF.")
    document.extracted_text = text
    await session.commit()

    if not llm.configured:
        return await _fail(session, document, "Text was extracted, but AI analysis is not configured on this server.", keep_text=True)
    try:
        extraction = await llm.complete_structured(
            system=prompts.DOCUMENT_EXTRACTION,
            messages=[{"role": "user", "content": prompts.tagged("document", text[:AI_CHARS])}],
            schema=DocumentExtraction,
            max_tokens=4000,
        )
    except AppError as exc:
        return await _fail(session, document, f"Text was extracted, but analysis failed: {exc.message}", keep_text=True)

    lowered = text.lower()
    for fact in extraction.facts:
        # Links must appear verbatim in the document; other facts are kept for user review.
        if fact.kind == "link" and fact.value.lower() not in lowered:
            continue
        session.add(ProfileFact(
            user_id=document.user_id, kind=fact.kind, value=fact.value.strip()[:500], detail=fact.detail, source="document",
            document_id=document.id, confirmed=False,
        ))
    document.processing_status = "processed"
    document.error = None
    document.processed_at = datetime.now(UTC)
    await session.commit()
    return document


async def retry(session: AsyncSession, llm: LLMClient, user_id: UUID, document_id: UUID) -> Document:
    document = await session.scalar(
        select(Document)
        .options(undefer(Document.content))
        .where(Document.id == document_id, Document.user_id == user_id)
    )
    if document is None:
        raise NotFoundError("Document not found")
    if document.processing_status not in ("failed", "processing"):
        raise BadRequestError("Only a document that did not finish can be read again.", code="document_not_retryable")
    # Failed OCR/AI attempts may have left suggestions; do not duplicate them.
    await session.execute(delete(ProfileFact).where(ProfileFact.user_id == user_id, ProfileFact.document_id == document.id, ProfileFact.confirmed.is_(False)))
    document.processing_status = "processing"
    document.error = None
    document.processed_at = None
    await session.commit()
    return await _process(session, llm, document)


async def _fail(session: AsyncSession, document: Document, message: str, *, keep_text: bool = False) -> Document:
    document.processing_status = "failed"
    document.error = message
    document.processed_at = datetime.now(UTC)
    if not keep_text:
        document.extracted_text = None
    await session.commit()
    return document


async def list_documents(session: AsyncSession, user_id: UUID) -> list[Document]:
    return list((await session.scalars(
        select(Document).where(Document.user_id == user_id).order_by(Document.created_at.desc())
    )).all())


async def get_document(session: AsyncSession, user_id: UUID, document_id: UUID, *, with_text: bool = False) -> Document:
    query = select(Document).where(Document.id == document_id, Document.user_id == user_id)
    if with_text:
        query = query.options(undefer(Document.extracted_text))
    document = await session.scalar(query)
    if document is None:
        raise NotFoundError("Document not found")
    return document


async def document_facts(session: AsyncSession, user_id: UUID, document_id: UUID) -> list[ProfileFact]:
    return list((await session.scalars(
        select(ProfileFact).where(ProfileFact.user_id == user_id, ProfileFact.document_id == document_id).order_by(ProfileFact.created_at)
    )).all())


async def delete_document(session: AsyncSession, user_id: UUID, document_id: UUID, *, keep_confirmed_facts: bool = True) -> None:
    """Delete the file, its extracted text and unconfirmed facts. Facts the
    user already confirmed into their profile are kept unless asked otherwise."""
    document = await get_document(session, user_id, document_id)
    facts = delete(ProfileFact).where(ProfileFact.user_id == user_id, ProfileFact.document_id == document_id)
    if keep_confirmed_facts:
        facts = facts.where(ProfileFact.confirmed.is_(False))
    await session.execute(facts)
    await session.delete(document)
    await session.commit()


async def search(session: AsyncSession, user_id: UUID, query: str, limit: int = 3) -> list[dict]:
    """Full-text search over the user's own documents (used by chat)."""
    tsquery = func.websearch_to_tsquery("english", query)
    rows = (await session.execute(
        select(Document.filename, func.ts_headline(
            "english", Document.extracted_text, tsquery, "MaxFragments=3, MaxWords=40, MinWords=15",
        ))
        .where(Document.user_id == user_id, Document.processing_status.in_(["processed", "failed"]))
        .where(Document.extracted_text.is_not(None))
        .where(literal_column("documents.search_vector").op("@@")(tsquery))
        .limit(limit)
    )).all()
    return [{"filename": filename, "excerpt": excerpt} for filename, excerpt in rows]
